"""
What Stage 15 adds
Create Mode button.
Click directly on the QC map to collect X/Y coordinates.
Live preview of the clicked points and geometry.
Displays the number of collected points.
Uses the currently selected geometry type:
Points → each click becomes a point feature.
Line (LineString) → clicks form one line.
Polygon → clicks form a polygon boundary.
Finish & Generate opens the normal shapefile save dialog.
Cancel Create discards the current construction.
Coordinate system is locked while creating so it cannot accidentally change.
Existing feature inspection remains available when Create Mode is off.
Polygon validation from Stage 14 is retained, including self-intersection/zero-area checks.
Created shapefile is automatically added to Active Layers and displayed on the QC canvas.

I also syntax-checked the Stage 15 application together with the two refactored support modules successfully.

Suggested test sequence
Load a shapefile so you have a visible map.
Select Points → Create Mode → click several locations → Finish & Generate.
Repeat with Line.
Repeat with Polygon.
Test polygon with only 2 clicks — should be rejected.
Test a bow-tie/self-intersecting polygon — should be rejected.
Confirm the newly created shapefile appears in Active Layers.
Turn Create Mode off and confirm normal feature inspection still works.
"""

import sys
import time
from datetime import datetime
from html import escape
from pathlib import Path
import numpy as np
import geopandas as gpd
import shapely
from shapely.geometry import Point, LineString, Polygon, box
import pyqtgraph as pg
from pyproj import CRS

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPlainTextEdit, QComboBox, QPushButton, QFileDialog, QLabel,
    QSplitter, QListWidget, QListWidgetItem, QMessageBox,
    QColorDialog, QMenu, QToolBar, QDialog, QFormLayout, QTextEdit,
    QTableView, QAbstractItemView, QLineEdit, QHeaderView
)
from PySide6.QtCore import (
    Qt, Signal, QThreadPool, QTimer, QSettings
)
from PySide6.QtGui import QPixmap, QColor, QIcon, QAction, QShortcut

# Global pyqtgraph settings
pg.setConfigOption('background', '#ffffff')
pg.setConfigOption('foreground', '#202020')
pg.setConfigOptions(antialias=True, useOpenGL=True)  # OpenGL needs PyOpenGL installed


def bouquet_studio_stylesheet() -> str:
    return """
        QMainWindow { background-color: #f2f2f2; }
        QWidget { color: #202020; font-family: Segoe UI, sans-serif; }
        QPlainTextEdit { background-color: #ffffff; border: 1px solid #777777; padding: 5px; }
        QPushButton { background-color: #ffffff; border: 1px solid #777777; padding: 6px; border-radius: 3px; }
        QPushButton:hover { border: 1px solid cyan; background-color: #f8f8f8; }
        QComboBox { background-color: #ffffff; border: 1px solid #777777; padding: 4px; }
        QListWidget { background-color: #ffffff; border: 1px solid #777777; padding: 3px; }
        QListWidget::item { padding: 4px; }
        QListWidget::item:hover { background-color: #f0f0f0; }
        QStatusBar { background-color: #e0e0e0; color: #202020; padding-left: 5px; border-top: 1px solid #777777; }
        QMessageBox { background-color: #f2f2f2; }
        QMenu { background-color: #ffffff; border: 1px solid #777777; }
        QMenu::item:selected { background-color: #e0e0e0; }
    """


def make_color_icon(color_hex: str) -> QIcon:
    pixmap = QPixmap(12, 12)
    pixmap.fill(QColor(color_hex))
    return QIcon(pixmap)


def make_status_icon(status: str) -> QIcon:
    """Create a compact QC status indicator for the Active Layers list."""
    status = status.upper()
    color = {
        "PASS": "#2e7d32",
        "WARNING": "#f9a825",
        "ERROR": "#c62828",
        "CHECKING...": "#757575",
    }.get(status, "#757575")
    pixmap = QPixmap(10, 10)
    pixmap.fill(QColor(color))
    return QIcon(pixmap)


from pipeline_route_creator_workers import (
    _LODWorker,
    _FeatureIndexWorker,
    _GeometryQCWorker,
    _build_gis_render_arrays,
)

class QCMapCanvas(pg.PlotWidget):
    status_message = Signal(str)
    layer_added = Signal(str, str, str)  # (layer_id, display_name, color_hex)
    layers_cleared = Signal()
    layer_qc_updated = Signal(str, str)  # (layer_id, status)
    create_points_changed = Signal(int)
    create_finished = Signal(object)

    PALETTE = ['#00cccc', '#e6007e', '#d9a100', '#2e7d32', '#6a1b9a', '#d84315']

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.showGrid(x=True, y=True, alpha=0.3)
        self.setAspectLocked(True)

        self.getAxis('bottom').setStyle(tickTextOffset=8)
        self.getAxis('left').setStyle(tickTextOffset=8)

        self.loaded_layers = {}  # layer_id -> layer metadata
        self.color_index = 0

        # Stage 3: cached multi-resolution rendering.
        # Geometry simplification happens in a worker, never during pan/zoom.
        self._lod_pool = QThreadPool(self)
        self._lod_pool.setMaxThreadCount(1)
        self._qc_generation = 0

        self._lod_switch_timer = QTimer(self)
        self._lod_switch_timer.setSingleShot(True)
        self._lod_switch_timer.timeout.connect(self._switch_visible_lod)

        self._lod_generation = 0

        self.getViewBox().sigRangeChanged.connect(
            self._schedule_lod_switch
        )

        # Stage 4: navigation and cursor coordinate readout.
        self.scene().sigMouseMoved.connect(self._mouse_moved)
        self._last_mouse_status_time = 0.0
        self._last_mouse_status_text = ""

        # Stage 7: feature inspection.
        self.scene().sigMouseClicked.connect(self._mouse_clicked)
        self._selected_feature_item = None
        self._selected_feature = None

        # Stage 15: interactive shapefile creation.
        self._create_mode = False
        self._create_geometry_type = "Points"
        self._create_points = []
        self._create_preview_items = []
        self._create_point_item = None
        self._create_line_item = None

    # ------------------------------------------------------------------
    # Stage 3: cached multi-resolution rendering
    # ------------------------------------------------------------------

    def _estimate_lod_tolerances(self, gdf):
        """
        Estimate three simplification distances from the layer extent.

        The exact values are deliberately data-scale based rather than
        hard-coded metre/degree values, so they work with projected and
        geographic CRSs.
        """
        try:
            min_x, min_y, max_x, max_y = gdf.total_bounds

            span = max(
                abs(float(max_x) - float(min_x)),
                abs(float(max_y) - float(min_y)),
            )

            if not np.isfinite(span) or span <= 0:
                span = 1.0

            # Full is already rendered. These are progressively lighter
            # cached representations.
            base = span / max(
                2000.0,
                float(max(
                    self.viewport().width(),
                    self.viewport().height(),
                )) * 2.0,
            )

            base = max(base, 1e-12)

            return {
                "medium": base * 2.0,
                "low": base * 8.0,
                "very_low": base * 32.0,
            }

        except Exception:
            return {
                "medium": 1.0,
                "low": 4.0,
                "very_low": 16.0,
            }

    def _start_lod_build(self, layer_id):
        layer = self.loaded_layers.get(layer_id)

        if not layer or not layer.get("has_complex_geometry"):
            return

        if layer.get("lod_build_started"):
            return

        gdf = layer.get("gdf")

        if gdf is None or gdf.empty:
            return

        layer["lod_build_started"] = True
        generation = self._lod_generation

        worker = _LODWorker(
            layer_id,
            generation,
            gdf,
            self._estimate_lod_tolerances(gdf),
        )

        worker.signals.finished.connect(
            self._lod_cache_ready
        )
        worker.signals.failed.connect(
            self._lod_cache_failed
        )

        layer["lod_worker"] = worker
        self._lod_pool.start(worker)

        self.status_message.emit(
            f"Preparing fast zoom cache for {layer['name']}..."
        )

    def _create_curve_item(self, x, y, color):
        if x is None or x.size == 0:
            return None

        curve = pg.PlotCurveItem(
            x,
            y,
            pen=pg.mkPen(
                color=color,
                width=1,
            ),
            connect="finite",
            antialias=False,
            skipFiniteCheck=False,
        )

        self.addItem(curve)
        return curve

    def _lod_cache_ready(self, layer_id, generation, arrays):
        if generation != self._lod_generation:
            return

        layer = self.loaded_layers.get(layer_id)

        if layer is None:
            return

        color = layer["color"]
        lod_items = layer.setdefault("lod_items", {})

        for level in ("medium", "low", "very_low"):
            data = arrays.get(level)

            if data is None:
                continue

            x, y = data
            item = self._create_curve_item(
                x,
                y,
                color,
            )

            if item is not None:
                item.setVisible(False)
                lod_items[level] = item

        layer["lod_cache_ready"] = True
        layer["lod_worker"] = None

        self._switch_visible_lod()

        self.status_message.emit(
            f"Fast zoom cache ready: {layer['name']}."
        )

    def _lod_cache_failed(self, layer_id, generation, message):
        if generation != self._lod_generation:
            return

        layer = self.loaded_layers.get(layer_id)

        if layer is not None:
            layer["lod_worker"] = None
            layer["lod_build_started"] = False

            self.status_message.emit(
                f"Fast zoom cache failed for "
                f"{layer['name']}: {message}"
            )

    def _current_lod_level(self, layer):
        """
        Cheap LOD selection.

        This is the only work performed on pan/zoom. No Shapely calls,
        coordinate extraction, or PlotCurveItem creation happen here.
        """
        if not layer.get("has_complex_geometry"):
            return "full"

        try:
            x_range, y_range = self.viewRange()

            view_span = max(
                abs(float(x_range[1]) - float(x_range[0])),
                abs(float(y_range[1]) - float(y_range[0])),
            )

            extent = layer.get("extent_span", 0.0)

            if extent <= 0 or view_span <= 0:
                return "full"

            ratio = view_span / extent

            # Close view -> full resolution.
            # Further away -> increasingly simplified cached data.
            if ratio <= 0.18:
                return "full"
            if ratio <= 0.55:
                return "medium"
            if ratio <= 1.50:
                return "low"
            return "very_low"

        except Exception:
            return "full"

    def _schedule_lod_switch(self, *args):
        if not self.loaded_layers:
            return

        # Debounce only the cheap visibility operation. There is no expensive
        # geometry work behind this timer.
        self._lod_switch_timer.start(30)

    def _switch_visible_lod(self):
        for layer in self.loaded_layers.values():

            if not layer.get("has_complex_geometry"):
                continue

            lod_items = layer.get("lod_items", {})

            desired = self._current_lod_level(layer)

            # Full-resolution item is the original item created at load time.
            candidates = {
                "full": layer.get("geometry_item"),
                "medium": lod_items.get("medium"),
                "low": lod_items.get("low"),
                "very_low": lod_items.get("very_low"),
            }

            selected = candidates.get(desired)

            # If the requested cache has not arrived yet, choose the best
            # available representation without rebuilding anything.
            if selected is None:
                for fallback in (
                    "full",
                    "medium",
                    "low",
                    "very_low",
                ):
                    if candidates.get(fallback) is not None:
                        selected = candidates[fallback]
                        break

            for item in candidates.values():
                if item is not None:
                    item.setVisible(
                        item is selected
                        and layer.get("visible", True)
                    )

            layer["active_lod"] = desired

    def _remove_layer_graphics(self, layer):
        items = set()

        for item in layer.get("items", []):
            if item is not None:
                items.add(item)

        geometry_item = layer.get("geometry_item")
        if geometry_item is not None:
            items.add(geometry_item)

        for item in layer.get("lod_items", {}).values():
            if item is not None:
                items.add(item)

        for item in items:
            try:
                self.removeItem(item)
            except Exception:
                pass

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.accept()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.accept()
        else:
            event.ignore()

    def dropEvent(self, event):
        urls = event.mimeData().urls()
        shp_files = [url.toLocalFile() for url in urls if url.toLocalFile().lower().endswith('.shp')]

        if not shp_files:
            self.status_message.emit("No valid .shp files found in drop payload.")
            return

        loaded_count = 0
        for filepath in shp_files:
            if self.add_shapefile_layer(filepath):
                loaded_count += 1

        self.status_message.emit(f"Added {loaded_count} layer(s). Active layers: {len(self.loaded_layers)}.")

    @staticmethod
    def _build_layer_statistics(gdf, filepath: str):
        """Build lightweight statistics once, when a layer is loaded."""
        stats = {
            "file_size": 0,
            "feature_count": int(len(gdf)),
            "geometry_types": {},
            "null_count": 0,
            "empty_count": 0,
            "vertex_count": 0,
            "crs": str(gdf.crs) if gdf.crs is not None else "Not defined",
            "bounds": None,
        }

        try:
            stats["file_size"] = int(Path(filepath).stat().st_size)
        except Exception:
            pass

        try:
            types = gdf.geometry.geom_type.fillna("None").value_counts()
            stats["geometry_types"] = {
                str(name): int(count) for name, count in types.items()
            }
        except Exception:
            pass

        try:
            geom = gdf.geometry
            stats["null_count"] = int(geom.isna().sum())
            stats["empty_count"] = int(geom.notna().sum() and geom[geom.notna()].is_empty.sum())
        except Exception:
            pass

        try:
            # Shapely 2 provides a vectorized coordinate-count operation.
            # This avoids a Python loop over every feature during loading.
            coords_count = shapely.get_num_coordinates(gdf.geometry.array)
            stats["vertex_count"] = int(np.asarray(coords_count).sum())
        except Exception:
            # Compatibility fallback for older Shapely versions.
            try:
                stats["vertex_count"] = int(sum(
                    len(shapely.get_coordinates(geom))
                    for geom in gdf.geometry
                    if geom is not None and not geom.is_empty
                ))
            except Exception:
                pass

        try:
            bounds = np.asarray(gdf.total_bounds, dtype=float)
            if bounds.size == 4 and np.all(np.isfinite(bounds)):
                stats["bounds"] = tuple(float(v) for v in bounds)
        except Exception:
            pass

        return stats

    def get_layer_statistics(self, layer_id: str):
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return None
        return layer.get("statistics", {})

    def add_shapefile_layer(self, filepath: str, target_crs: str = None) -> bool:
        path_obj = Path(filepath)
        layer_id = str(path_obj.resolve())

        if layer_id in self.loaded_layers:
            self.status_message.emit(f"Layer '{path_obj.name}' is already loaded.")
            return False

        if target_crs is None and hasattr(self.window(), 'get_active_crs'):
            target_crs = self.window().get_active_crs()

        try:
            load_start = time.perf_counter()

            # pyogrio backend is significantly faster than fiona for large files.
            # Falls back automatically if pyogrio isn't installed.
            try:
                gdf = gpd.read_file(filepath, engine="pyogrio")
            except Exception:
                gdf = gpd.read_file(filepath)

            if gdf.empty:
                return False

            # --- CRS Validation & Conversion Prompt ---
            if target_crs:
                if gdf.crs is not None:
                    file_crs_obj = CRS.from_user_input(gdf.crs)
                    target_crs_obj = CRS.from_user_input(target_crs)

                    if file_crs_obj != target_crs_obj:
                        reply = QMessageBox.question(
                            self,
                            "CRS Mismatch Detected",
                            f"Layer '{path_obj.name}' CRS is:\n  • {file_crs_obj.name}\n\n"
                            f"Active Application CRS is:\n  • {target_crs_obj.name} ({target_crs})\n\n"
                            f"Would you like to reproject this layer to match the active system?",
                            QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel,
                            QMessageBox.Yes
                        )

                        if reply == QMessageBox.Yes:
                            gdf = gdf.to_crs(target_crs)
                            self.status_message.emit(f"Reprojected {path_obj.name} to {target_crs}.")
                        elif reply == QMessageBox.Cancel:
                            self.status_message.emit(f"Cancelled loading {path_obj.name}.")
                            return False
                else:
                    reply = QMessageBox.warning(
                        self,
                        "Missing CRS (.prj file)",
                        f"The file '{path_obj.name}' has no projection metadata (.prj file missing).\n\n"
                        f"Assume active CRS ({target_crs})?",
                        QMessageBox.Yes | QMessageBox.No,
                        QMessageBox.Yes
                    )
                    if reply == QMessageBox.Yes:
                        gdf.set_crs(target_crs, inplace=True)
                    else:
                        return False

            color = self.PALETTE[self.color_index % len(self.PALETTE)]
            self.color_index += 1

            graphic_items = []

            # 1. VECTORIZED POINT RENDERING
            point_mask = gdf.geometry.type.isin(['Point', 'MultiPoint'])
            if point_mask.any():
                pts = gdf.geometry[point_mask].explode(index_parts=False)
                coords = shapely.get_coordinates(pts.values)

                scatter = pg.ScatterPlotItem(
                    x=coords[:, 0],
                    y=coords[:, 1],
                    size=6,
                    pen=pg.mkPen(color=color, width=1),
                    brush=pg.mkBrush(color=color),
                    name=path_obj.name
                )
                self.addItem(scatter)
                graphic_items.append(scatter)

            # 2. LINE & POLYGON RENDERING
            #
            # Stage 3 keeps the original full-resolution item for immediate
            # display. Additional resolutions are generated once in the
            # background and cached as separate PlotCurveItems.
            line_mask = gdf.geometry.geom_type.isin(
                ['LineString', 'MultiLineString']
            )
            poly_mask = gdf.geometry.geom_type.isin(
                ['Polygon', 'MultiPolygon']
            )

            has_complex_geometry = bool(
                line_mask.any() or poly_mask.any()
            )

            geometry_item = None

            if has_complex_geometry:
                all_x, all_y = _build_gis_render_arrays(
                    gdf,
                    tolerance=0.0,
                )

                if all_x.size:
                    geometry_item = pg.PlotCurveItem(
                        all_x,
                        all_y,
                        pen=pg.mkPen(
                            color=color,
                            width=1,
                        ),
                        connect='finite',
                        antialias=False,
                        skipFiniteCheck=False,
                    )

                    self.addItem(geometry_item)
                    graphic_items.append(geometry_item)

            try:
                bounds = gdf.total_bounds
                extent_span = max(
                    abs(float(bounds[2]) - float(bounds[0])),
                    abs(float(bounds[3]) - float(bounds[1])),
                )
            except Exception:
                extent_span = 0.0

            elapsed = time.perf_counter() - load_start
            statistics = self._build_layer_statistics(gdf, filepath)

            self.loaded_layers[layer_id] = {
                'name': path_obj.name,
                'items': graphic_items,
                'geometry_item': geometry_item,
                'gdf': gdf if has_complex_geometry else None,
                'full_gdf': gdf,
                'has_complex_geometry': has_complex_geometry,
                'visible': True,
                'color': color,
                'extent_span': extent_span,
                'lod_items': {},
                'lod_cache_ready': False,
                'lod_build_started': False,
                'lod_worker': None,
                'active_lod': 'full',
                'statistics': statistics,
                'qc_status': 'CHECKING...',
                'qc_result': None,
                'qc_started': False,
                'qc_worker': None,
                'feature_index': None,
                'feature_index_started': False,
                'feature_index_worker': None,
            }

            self.status_message.emit(
                f"Loaded {path_obj.name}: "
                f"{statistics['feature_count']:,} feature(s), "
                f"{statistics['vertex_count']:,} vertex(es), "
                f"{elapsed:.2f}s."
            )

            self.layer_added.emit(layer_id, path_obj.name, color)

            if has_complex_geometry:
                # The Active Layers entry is created first. The expensive
                # simplified representations are then prepared off the GUI
                # thread.
                self._start_lod_build(layer_id)

            # Geometry QC is also performed in the background so large files
            # remain interactive while checks are running.
            self._start_geometry_qc(layer_id)
            self._start_feature_index(layer_id)

            return True

        except Exception as e:
            self.status_message.emit(f"Error loading {path_obj.name}: {str(e)}")
            return False

    def _start_geometry_qc(self, layer_id):
        layer = self.loaded_layers.get(layer_id)
        if not layer or layer.get("qc_started"):
            return

        gdf = layer.get("gdf")
        if gdf is None:
            # Point-only layers still need QC, so use the full GeoDataFrame.
            gdf = layer.get("full_gdf")
        if gdf is None or gdf.empty:
            return

        layer["qc_started"] = True
        generation = self._qc_generation
        worker = _GeometryQCWorker(layer_id, generation, gdf)
        worker.signals.finished.connect(self._geometry_qc_ready)
        worker.signals.failed.connect(self._geometry_qc_failed)
        layer["qc_worker"] = worker
        layer["qc_status"] = "CHECKING..."
        self._lod_pool.start(worker)

    def _geometry_qc_ready(self, layer_id, generation, result):
        if generation != self._qc_generation:
            return
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return

        layer["qc_result"] = result
        layer["qc_status"] = result.get("status", "WARNING")
        layer["qc_worker"] = None
        self.layer_qc_updated.emit(layer_id, layer["qc_status"])

        status = layer["qc_status"]
        if result.get("issues"):
            self.status_message.emit(
                f"QC {status}: {layer['name']} — "
                + "; ".join(result["issues"][:3])
            )
        else:
            self.status_message.emit(f"QC PASS: {layer['name']}.")

    def _geometry_qc_failed(self, layer_id, generation, message):
        if generation != self._qc_generation:
            return
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return
        layer["qc_worker"] = None
        layer["qc_started"] = False
        layer["qc_status"] = "FAILED"
        layer["qc_result"] = {"status": "FAILED", "issues": [message]}
        self.layer_qc_updated.emit(layer_id, "FAILED")
        self.status_message.emit(f"Geometry QC failed for {layer['name']}: {message}")

    def _start_feature_index(self, layer_id):
        layer = self.loaded_layers.get(layer_id)
        if not layer or layer.get("feature_index_started"):
            return

        gdf = layer.get("full_gdf")
        if gdf is None or gdf.empty:
            return

        layer["feature_index_started"] = True
        worker = _FeatureIndexWorker(layer_id, self._qc_generation, gdf)
        worker.signals.finished.connect(self._feature_index_ready)
        worker.signals.failed.connect(self._feature_index_failed)
        layer["feature_index_worker"] = worker
        self._lod_pool.start(worker)

    def _feature_index_ready(self, layer_id, generation, index):
        if generation != self._qc_generation:
            return
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return
        layer["feature_index"] = index
        layer["feature_index_worker"] = None
        self.status_message.emit(f"Feature inspection ready: {layer['name']}.")

    def _feature_index_failed(self, layer_id, generation, message):
        if generation != self._qc_generation:
            return
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return
        layer["feature_index_worker"] = None
        layer["feature_index_started"] = False
        self.status_message.emit(
            f"Feature inspection index unavailable for {layer['name']}: {message}"
        )

    def _pick_feature(self, scene_pos):
        """Return (layer_id, positional_index, distance) for a map click."""
        try:
            view_point = self.getViewBox().mapSceneToView(scene_pos)
            px = float(view_point.x())
            py = float(view_point.y())
            if not (np.isfinite(px) and np.isfinite(py)):
                return None

            x_range, y_range = self.viewRange()
            view_width = abs(float(x_range[1]) - float(x_range[0]))
            pixels = max(float(self.viewport().width()), 1.0)
            tolerance = max(view_width / pixels * 8.0, 1e-12)
            click_point = Point(px, py)
            search_area = box(px - tolerance, py - tolerance,
                              px + tolerance, py + tolerance)

            best = None
            for layer_id, layer in self.loaded_layers.items():
                if not layer.get("visible", True):
                    continue
                gdf = layer.get("full_gdf")
                index = layer.get("feature_index")
                if gdf is None:
                    continue
                if index is None:
                    continue

                try:
                    candidates = index.query(search_area, predicate="intersects")
                except TypeError:
                    candidates = index.query(search_area)

                for pos in np.asarray(candidates, dtype=int):
                    try:
                        geom = gdf.geometry.iloc[int(pos)]
                    except Exception:
                        continue
                    if geom is None or geom.is_empty:
                        continue

                    try:
                        distance = float(geom.distance(click_point))
                    except Exception:
                        continue

                    if distance <= tolerance:
                        # Prefer a containing polygon over a nearby boundary.
                        contains_bonus = 0 if geom.geom_type in ("Polygon", "MultiPolygon") and geom.contains(click_point) else 1
                        score = (contains_bonus, distance)
                        if best is None or score < best[0]:
                            best = (score, layer_id, int(pos), distance)

            if best is None:
                return None
            return best[1], best[2], best[3]
        except Exception:
            return None

    def _clear_feature_selection(self):
        if self._selected_feature_item is not None:
            try:
                self.removeItem(self._selected_feature_item)
            except Exception:
                pass
        self._selected_feature_item = None
        self._selected_feature = None

    def _highlight_feature(self, layer_id, position):
        self._clear_feature_selection()
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return
        gdf = layer.get("full_gdf")
        if gdf is None:
            return
        try:
            geom = gdf.geometry.iloc[position]
            if geom is None or geom.is_empty:
                return

            if geom.geom_type in ("Point", "MultiPoint"):
                coords = shapely.get_coordinates(geom)
                item = pg.ScatterPlotItem(
                    x=coords[:, 0],
                    y=coords[:, 1],
                    size=12,
                    pen=pg.mkPen(color="#ff0000", width=2),
                    brush=None,
                )
            else:
                coords_x, coords_y = _build_single_geometry_render_arrays(geom)
                if not coords_x.size:
                    return
                item = pg.PlotCurveItem(
                    coords_x, coords_y,
                    pen=pg.mkPen(color="#ff0000", width=3),
                    connect="finite",
                    antialias=False,
                    skipFiniteCheck=False,
                )

            self.addItem(item)
            self._selected_feature_item = item
            self._selected_feature = (layer_id, position)
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Stage 15: interactive Create Mode
    # ------------------------------------------------------------------

    def set_create_mode(self, enabled, geometry_type=None):
        enabled = bool(enabled)
        if geometry_type is not None:
            self._create_geometry_type = geometry_type

        if enabled == self._create_mode:
            return

        self._create_mode = enabled
        if enabled:
            self._clear_feature_selection()
            self._create_points = []
            self._clear_create_preview()
            self.create_points_changed.emit(0)
            self.status_message.emit(
                "Create Mode: click points on the map, then choose Finish & Generate."
            )
        else:
            self._create_points = []
            self._clear_create_preview()
            self.create_points_changed.emit(0)

    def cancel_create_mode(self):
        if not self._create_mode:
            return
        self._create_mode = False
        self._create_points = []
        self._clear_create_preview()
        self.create_points_changed.emit(0)
        self.status_message.emit("Create Mode cancelled.")

    def creation_points(self):
        return list(self._create_points)

    def finish_create_mode(self):
        if not self._create_mode:
            return
        points = list(self._create_points)
        self._create_mode = False
        self._clear_create_preview()
        self.create_points_changed.emit(len(points))
        self.create_finished.emit(points)

    def _clear_create_preview(self):
        for item in self._create_preview_items:
            try:
                self.removeItem(item)
            except Exception:
                pass
        self._create_preview_items = []
        self._create_point_item = None
        self._create_line_item = None

    def _update_create_preview(self):
        self._clear_create_preview()
        if not self._create_points:
            return

        coords = np.asarray(self._create_points, dtype=float)
        self._create_point_item = pg.ScatterPlotItem(
            x=coords[:, 0],
            y=coords[:, 1],
            size=10,
            pen=pg.mkPen(color="#1565c0", width=2),
            brush=pg.mkBrush("#ffffff"),
        )
        self.addItem(self._create_point_item)
        self._create_preview_items.append(self._create_point_item)

        if len(self._create_points) >= 2:
            line_x = coords[:, 0]
            line_y = coords[:, 1]
            if self._create_geometry_type == "Polygon" and len(self._create_points) >= 3:
                line_x = np.r_[line_x, line_x[0]]
                line_y = np.r_[line_y, line_y[0]]

            self._create_line_item = pg.PlotCurveItem(
                line_x,
                line_y,
                pen=pg.mkPen(color="#1565c0", width=2),
                connect="all",
                antialias=False,
            )
            self.addItem(self._create_line_item)
            self._create_preview_items.append(self._create_line_item)

    def _mouse_clicked(self, event):
        if event.button() != Qt.LeftButton:
            return
        scene_pos = event.scenePos()
        if not self.sceneBoundingRect().contains(scene_pos):
            return

        if self._create_mode:
            view_pos = self.getViewBox().mapSceneToView(scene_pos)
            x = float(view_pos.x())
            y = float(view_pos.y())
            if np.isfinite(x) and np.isfinite(y):
                self._create_points.append((x, y))
                self._update_create_preview()
                self.create_points_changed.emit(len(self._create_points))
                self.status_message.emit(
                    f"Create Mode: point {len(self._create_points)} — "
                    f"X {x:,.3f}, Y {y:,.3f}."
                )
            return

        picked = self._pick_feature(scene_pos)
        if picked is None:
            self._clear_feature_selection()
            return

        layer_id, position, distance = picked
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return

        self._highlight_feature(layer_id, position)
        self.show_feature_information(layer_id, position, distance)

    def show_feature_information(self, layer_id, position, distance=None):
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return
        gdf = layer.get("full_gdf")
        if gdf is None:
            return

        try:
            row = gdf.iloc[position]
            geometry = row.geometry
            original_index = row.name
        except Exception:
            return

        dialog = QDialog(self)
        dialog.setWindowTitle(f"Feature Inspection — {layer.get('name', 'Layer')}")
        dialog.resize(620, 520)
        layout = QVBoxLayout(dialog)

        title = QLabel(
            f"Layer: {layer.get('name', '')}    |    Feature: {position + 1:,}"
        )
        layout.addWidget(title)

        text = QTextEdit()
        text.setReadOnly(True)
        lines = [
            f"Feature position: {position + 1:,}",
            f"Original index: {original_index}",
            f"Geometry type: {geometry.geom_type if geometry is not None else 'None'}",
        ]
        if geometry is not None and not geometry.is_empty:
            try:
                lines.append(f"Geometry length: {float(geometry.length):,.3f}")
            except Exception:
                pass
            try:
                lines.append(f"Geometry area: {float(geometry.area):,.3f}")
            except Exception:
                pass
        if distance is not None:
            lines.append(f"Click distance: {distance:,.6f}")

        lines.extend(["", "Attributes", "------------------------------"])
        for column in gdf.columns:
            if column == gdf.geometry.name:
                continue
            value = row[column]
            if isinstance(value, float):
                value_text = f"{value:,.6g}"
            else:
                value_text = str(value)
            lines.append(f"{column}: {value_text}")

        text.setPlainText("\n".join(lines))
        layout.addWidget(text)

        close_button = QPushButton("Close")
        close_button.clicked.connect(dialog.accept)
        layout.addWidget(close_button)
        dialog.exec()

    def zoom_to_feature(self, layer_id, position):
        """Highlight a feature and zoom the map to its bounds."""
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return
        gdf = layer.get("full_gdf")
        if gdf is None or position < 0 or position >= len(gdf):
            return

        try:
            geometry = gdf.geometry.iloc[position]
            if geometry is None or geometry.is_empty:
                return
            bounds = geometry.bounds
            if len(bounds) != 4 or not np.all(np.isfinite(bounds)):
                return
            xmin, ymin, xmax, ymax = map(float, bounds)
            if xmax <= xmin:
                xmax = xmin + 1.0
            if ymax <= ymin:
                ymax = ymin + 1.0

            self._highlight_feature(layer_id, position)
            self.setRange(
                xRange=(xmin, xmax),
                yRange=(ymin, ymax),
                padding=0.15,
            )
            self.status_message.emit(
                f"Feature {position + 1:,} selected in {layer['name']}."
            )
        except Exception as exc:
            self.status_message.emit(f"Unable to zoom to feature: {exc}")

    def show_attribute_table(self, layer_id: str):
        """Show a read-only, searchable attribute table for one layer."""
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return

        gdf = layer.get("full_gdf")
        if gdf is None:
            return

        geometry_column = gdf.geometry.name
        dialog = QDialog(self)
        dialog.setWindowTitle(f"Attribute Table — {layer.get('name', 'Layer')}")
        dialog.resize(900, 600)

        layout = QVBoxLayout(dialog)
        header = QLabel(
            f"{layer.get('name', '')}    |    {len(gdf):,} features"
        )
        layout.addWidget(header)

        search = QLineEdit()
        search.setPlaceholderText("Filter attributes... (searches all columns)")
        layout.addWidget(search)

        model = _AttributeTableModel(gdf, geometry_column, dialog)
        proxy = _AttributeFilterProxyModel(dialog)
        proxy.setSourceModel(model)
        proxy.setFilterCaseSensitivity(Qt.CaseInsensitive)
        proxy.setSortCaseSensitivity(Qt.CaseInsensitive)
        proxy.setDynamicSortFilter(True)

        table = QTableView()
        table.setModel(proxy)
        table.setSortingEnabled(True)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setAlternatingRowColors(True)
        table.setWordWrap(False)
        table.verticalHeader().setDefaultSectionSize(22)
        table.horizontalHeader().setStretchLastSection(True)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        table.setColumnWidth(0, 80)
        layout.addWidget(table)

        footer = QLabel(
            "Double-click a row to select and zoom to the feature."
        )
        layout.addWidget(footer)

        close_button = QPushButton("Close")
        close_button.clicked.connect(dialog.accept)
        layout.addWidget(close_button)

        def apply_filter(text):
            proxy.setFilterFixedString(text)

        search.textChanged.connect(apply_filter)

        def activate_row(proxy_index):
            if not proxy_index.isValid():
                return
            source_index = proxy.mapToSource(proxy_index)
            position = source_index.row()
            self.zoom_to_feature(layer_id, position)

        table.doubleClicked.connect(activate_row)
        dialog.exec()

    def get_geometry_qc(self, layer_id: str):
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return None
        return layer.get("qc_result")

    def show_geometry_qc(self, layer_id: str):
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return

        result = layer.get("qc_result")
        dialog = QDialog(self)
        dialog.setWindowTitle(f"Geometry QC — {layer.get('name', 'Layer')}")
        dialog.resize(620, 460)
        layout = QVBoxLayout(dialog)

        status = layer.get("qc_status", "CHECKING...")
        status_label = QLabel(f"QC STATUS: {status}")
        layout.addWidget(status_label)

        if result is None:
            text = "Geometry QC is still running in the background.\n\nClose this dialog and open it again when the status changes."
        else:
            lines = [
                f"Features:            {result.get('feature_count', 0):,}",
                f"Valid geometries:    {result.get('valid_count', 0):,}",
                f"Invalid geometries:  {result.get('invalid_count', 0):,}",
                f"Null geometries:     {result.get('null_count', 0):,}",
                f"Empty geometries:    {result.get('empty_count', 0):,}",
                f"Zero-length lines:   {result.get('zero_length_count', 0):,}",
                f"Duplicate geometries: {result.get('duplicate_count', 0):,}",
                "",
                "Geometry types:",
            ]
            for name, count in result.get("geometry_types", {}).items():
                lines.append(f"  {name}: {count:,}")

            lines.extend(["", "Issues:"])
            issues = result.get("issues", [])
            lines.extend(f"  • {issue}" for issue in issues) if issues else lines.append("  • None")
            text = "\n".join(lines)

        text_edit = QTextEdit()
        text_edit.setReadOnly(True)
        text_edit.setPlainText(text)
        layout.addWidget(text_edit)

        close_btn = QPushButton("Close")
        close_btn.clicked.connect(dialog.accept)
        layout.addWidget(close_btn)
        dialog.exec()

    def show_layer_information(self, layer_id: str):
        """Show lightweight, load-time statistics for one layer."""
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return

        stats = layer.get("statistics", {})
        dialog = QDialog(self)
        dialog.setWindowTitle(f"Layer Information — {layer.get('name', 'Layer')}")
        dialog.resize(560, 420)

        layout = QFormLayout(dialog)

        def add_row(label, value):
            layout.addRow(QLabel(str(label)), QLabel(str(value)))

        size_bytes = int(stats.get("file_size", 0) or 0)
        if size_bytes < 1024:
            size_text = f"{size_bytes:,} B"
        elif size_bytes < 1024 ** 2:
            size_text = f"{size_bytes / 1024:.1f} KB"
        elif size_bytes < 1024 ** 3:
            size_text = f"{size_bytes / (1024 ** 2):.1f} MB"
        else:
            size_text = f"{size_bytes / (1024 ** 3):.2f} GB"

        geometry_types = stats.get("geometry_types", {})
        geometry_text = ", ".join(
            f"{name}: {count:,}" for name, count in geometry_types.items()
        ) or "Not available"

        bounds = stats.get("bounds")
        if bounds:
            xmin, ymin, xmax, ymax = bounds
            bounds_text = (
                f"X: {xmin:,.3f} to {xmax:,.3f}\n"
                f"Y: {ymin:,.3f} to {ymax:,.3f}"
            )
        else:
            bounds_text = "Not available"

        add_row("File", layer.get("name", ""))
        add_row("File size", size_text)
        add_row("Features", f"{stats.get('feature_count', 0):,}")
        add_row("Vertices", f"{stats.get('vertex_count', 0):,}")
        add_row("Geometry types", geometry_text)
        add_row("Null geometries", f"{stats.get('null_count', 0):,}")
        add_row("Empty geometries", f"{stats.get('empty_count', 0):,}")
        add_row("CRS", stats.get("crs", "Not defined"))
        add_row("Bounds", bounds_text)

        close_button = QPushButton("Close")
        close_button.clicked.connect(dialog.accept)
        layout.addRow("", close_button)

        dialog.exec()

    def _report_rows(self):
        """Return a flat, report-friendly summary for all loaded layers."""
        rows = []
        for layer_id, layer in self.loaded_layers.items():
            stats = layer.get("statistics", {}) or {}
            qc = layer.get("qc_result") or {}
            status = layer.get("qc_status", "CHECKING...")
            geometry_types = stats.get("geometry_types", {}) or qc.get("geometry_types", {}) or {}
            geometry_text = ", ".join(
                f"{name}: {count:,}" for name, count in geometry_types.items()
            )
            bounds = stats.get("bounds")
            bounds_text = ""
            if bounds:
                bounds_text = ", ".join(f"{float(value):.6f}" for value in bounds)

            rows.append({
                "layer_id": layer_id,
                "file": layer.get("name", ""),
                "file_path": layer_id,
                "file_size_bytes": int(stats.get("file_size", 0) or 0),
                "feature_count": int(stats.get("feature_count", 0) or 0),
                "vertex_count": int(stats.get("vertex_count", 0) or 0),
                "geometry_types": geometry_text,
                "null_count": int(qc.get("null_count", stats.get("null_count", 0)) or 0),
                "empty_count": int(qc.get("empty_count", stats.get("empty_count", 0)) or 0),
                "valid_count": int(qc.get("valid_count", 0) or 0),
                "invalid_count": int(qc.get("invalid_count", 0) or 0),
                "zero_length_count": int(qc.get("zero_length_count", 0) or 0),
                "duplicate_count": int(qc.get("duplicate_count", 0) or 0),
                "crs": stats.get("crs", "Not defined"),
                "bounds": bounds_text,
                "status": status,
                "issues": "; ".join(qc.get("issues", []) or []),
            })
        return rows

    def _report_summary(self, rows):
        counts = {"PASS": 0, "WARNING": 0, "ERROR": 0, "CHECKING...": 0, "FAILED": 0}
        for row in rows:
            counts[row["status"]] = counts.get(row["status"], 0) + 1
        if counts.get("ERROR", 0):
            overall = "ERROR"
        elif counts.get("WARNING", 0) or counts.get("FAILED", 0):
            overall = "WARNING"
        elif counts.get("CHECKING...", 0):
            overall = "CHECKING"
        else:
            overall = "PASS"
        return overall, counts

    def _build_html_qc_report(self, rows, overall, counts):
        generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        total_features = sum(row["feature_count"] for row in rows)
        total_vertices = sum(row["vertex_count"] for row in rows)

        summary_cards = "".join(
            f"<div class='card'><div class='label'>{escape(label)}</div>"
            f"<div class='value'>{value:,}</div></div>"
            for label, value in [
                ("Layers", len(rows)),
                ("Features", total_features),
                ("Vertices", total_vertices),
                ("PASS", counts.get("PASS", 0)),
                ("WARNING", counts.get("WARNING", 0)),
                ("ERROR", counts.get("ERROR", 0)),
            ]
        )

        table_rows = []
        for row in rows:
            issue_text = escape(row["issues"]) if row["issues"] else "—"
            table_rows.append(
                "<tr>"
                f"<td>{escape(row['file'])}</td>"
                f"<td>{escape(row['status'])}</td>"
                f"<td>{row['feature_count']:,}</td>"
                f"<td>{row['vertex_count']:,}</td>"
                f"<td>{row['valid_count']:,}</td>"
                f"<td>{row['invalid_count']:,}</td>"
                f"<td>{row['null_count']:,}</td>"
                f"<td>{row['empty_count']:,}</td>"
                f"<td>{row['zero_length_count']:,}</td>"
                f"<td>{row['duplicate_count']:,}</td>"
                f"<td>{escape(str(row['crs']))}</td>"
                f"<td>{issue_text}</td>"
                "</tr>"
            )

        return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pipeline Route QC Report</title>
<style>
body {{ font-family: Segoe UI, Arial, sans-serif; margin: 32px; color: #202020; background: #f7f7f7; }}
h1 {{ margin-bottom: 4px; }}
.meta {{ color: #666; margin-bottom: 24px; }}
.summary {{ display: flex; flex-wrap: wrap; gap: 10px; margin-bottom: 24px; }}
.card {{ background: white; border: 1px solid #ccc; padding: 12px 18px; min-width: 100px; }}
.label {{ font-size: 12px; color: #666; }}
.value {{ font-size: 22px; font-weight: 600; margin-top: 3px; }}
.overall {{ font-size: 20px; font-weight: 600; margin: 12px 0 20px; }}
table {{ border-collapse: collapse; width: 100%; background: white; font-size: 12px; }}
th, td {{ border: 1px solid #ccc; padding: 7px; text-align: left; vertical-align: top; }}
th {{ background: #eaeaea; }}
.pass {{ font-weight: 600; }}
.warning {{ font-weight: 600; }}
.error {{ font-weight: 600; }}
</style>
</head>
<body>
<h1>Pipeline Route QC Report</h1>
<div class="meta">Generated: {escape(generated)}</div>
<div class="overall">Overall status: {escape(overall)}</div>
<div class="summary">{summary_cards}</div>
<table>
<thead><tr>
<th>Layer</th><th>Status</th><th>Features</th><th>Vertices</th>
<th>Valid</th><th>Invalid</th><th>Null</th><th>Empty</th>
<th>Zero-length</th><th>Duplicates</th><th>CRS</th><th>Issues</th>
</tr></thead>
<tbody>{''.join(table_rows)}</tbody>
</table>
</body>
</html>"""

    def _export_html_report(self, filename, rows, overall, counts):
        Path(filename).write_text(
            self._build_html_qc_report(rows, overall, counts),
            encoding="utf-8",
        )

    def _export_excel_report(self, filename, rows, overall, counts):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "QC Summary"

        ws["A1"] = "Pipeline Route QC Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = "Generated"
        ws["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws["A3"] = "Overall Status"
        ws["B3"] = overall

        headers = [
            "Layer", "Status", "File Size (bytes)", "Features", "Vertices",
            "Geometry Types", "Valid", "Invalid", "Null", "Empty",
            "Zero-length", "Duplicates", "CRS", "Bounds", "Issues", "File Path",
        ]
        start_row = 5
        for col, header in enumerate(headers, 1):
            cell = ws.cell(start_row, col, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EAEAEA")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        for row_number, row in enumerate(rows, start_row + 1):
            values = [
                row["file"], row["status"], row["file_size_bytes"], row["feature_count"],
                row["vertex_count"], row["geometry_types"], row["valid_count"],
                row["invalid_count"], row["null_count"], row["empty_count"],
                row["zero_length_count"], row["duplicate_count"], row["crs"],
                row["bounds"], row["issues"], row["file_path"],
            ]
            for col, value in enumerate(values, 1):
                ws.cell(row_number, col, value)
                ws.cell(row_number, col).alignment = Alignment(vertical="top", wrap_text=True)

        # Small status summary at the right of the report.
        summary_col = len(headers) + 2
        ws.cell(1, summary_col, "Status Counts").font = Font(bold=True)
        for offset, status in enumerate(("PASS", "WARNING", "ERROR", "CHECKING...", "FAILED"), 2):
            ws.cell(offset, summary_col, status)
            ws.cell(offset, summary_col + 1, counts.get(status, 0))

        ws.freeze_panes = "A6"
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
        for column_cells in ws.columns:
            max_length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = min(max(max_length, len(value)), 45)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(12, max_length + 2)

        wb.save(filename)

    def export_qc_report(self):
        """Export the current loaded-layer QC state to HTML and/or Excel."""
        rows = self._report_rows()
        if not rows:
            QMessageBox.information(self, "QC Report", "No layers are currently loaded.")
            return

        overall, counts = self._report_summary(rows)

        dialog = QDialog(self)
        dialog.setWindowTitle("Export QC Report")
        dialog.resize(520, 300)
        layout = QVBoxLayout(dialog)

        label = QLabel(
            f"{len(rows):,} layer(s) loaded\n"
            f"Overall status: {overall}\n\n"
            "Export a report containing the current Layer Information and Geometry QC results."
        )
        label.setWordWrap(True)
        layout.addWidget(label)

        html_btn = QPushButton("Export HTML Report...")
        excel_btn = QPushButton("Export Excel Report...")
        both_btn = QPushButton("Export Both...")
        close_btn = QPushButton("Close")
        layout.addWidget(html_btn)
        layout.addWidget(excel_btn)
        layout.addWidget(both_btn)
        layout.addWidget(close_btn)

        def export_html():
            filename, _ = QFileDialog.getSaveFileName(
                dialog, "Save HTML QC Report", "Pipeline_QC_Report.html",
                "HTML files (*.html);;All files (*)"
            )
            if not filename:
                return
            try:
                if not filename.lower().endswith(".html"):
                    filename += ".html"
                self._export_html_report(filename, rows, overall, counts)
                self.status_message.emit(f"QC HTML report exported: {Path(filename).name}")
                QMessageBox.information(dialog, "QC Report", f"HTML report saved to:\n{filename}")
            except Exception as exc:
                QMessageBox.critical(dialog, "Export Failed", str(exc))

        def export_excel():
            filename, _ = QFileDialog.getSaveFileName(
                dialog, "Save Excel QC Report", "Pipeline_QC_Report.xlsx",
                "Excel files (*.xlsx);;All files (*)"
            )
            if not filename:
                return
            try:
                if not filename.lower().endswith(".xlsx"):
                    filename += ".xlsx"
                self._export_excel_report(filename, rows, overall, counts)
                self.status_message.emit(f"QC Excel report exported: {Path(filename).name}")
                QMessageBox.information(dialog, "QC Report", f"Excel report saved to:\n{filename}")
            except Exception as exc:
                QMessageBox.critical(dialog, "Export Failed", str(exc))

        def export_both():
            directory = QFileDialog.getExistingDirectory(
                dialog, "Choose Report Folder"
            )
            if not directory:
                return
            try:
                html_path = Path(directory) / "Pipeline_QC_Report.html"
                excel_path = Path(directory) / "Pipeline_QC_Report.xlsx"
                self._export_html_report(str(html_path), rows, overall, counts)
                self._export_excel_report(str(excel_path), rows, overall, counts)
                self.status_message.emit("QC HTML and Excel reports exported.")
                QMessageBox.information(
                    dialog, "QC Report",
                    f"Reports saved to:\n{html_path}\n{excel_path}"
                )
            except Exception as exc:
                QMessageBox.critical(dialog, "Export Failed", str(exc))

        html_btn.clicked.connect(export_html)
        excel_btn.clicked.connect(export_excel)
        both_btn.clicked.connect(export_both)
        close_btn.clicked.connect(dialog.accept)
        dialog.exec()

    def change_layer_color(self, layer_id: str, new_color: str):
        layer = self.loaded_layers.get(layer_id)

        if layer is None:
            return

        layer["color"] = new_color

        items = list(layer.get("items", []))
        items.extend(
            item
            for item in layer.get("lod_items", {}).values()
            if item is not None
        )

        geometry_item = layer.get("geometry_item")
        if geometry_item is not None:
            items.append(geometry_item)

        for item in set(items):
            if isinstance(item, pg.ScatterPlotItem):
                item.setBrush(
                    pg.mkBrush(color=new_color)
                )
                item.setPen(
                    pg.mkPen(
                        color=new_color,
                        width=1,
                    )
                )

            elif isinstance(item, pg.PlotCurveItem):
                item.setPen(
                    pg.mkPen(
                        color=new_color,
                        width=1,
                    )
                )


    def remove_layer(self, layer_id: str):
        """Remove one layer and all of its cached graphics."""
        layer = self.loaded_layers.pop(layer_id, None)
        if layer is None:
            return

        if self._selected_feature is not None:
            selected_layer_id = self._selected_feature.get("layer_id")
            if selected_layer_id == layer_id:
                self._clear_feature_selection()

        self._remove_layer_graphics(layer)
        self._lod_generation += 1
        self._qc_generation += 1
        self._lod_switch_timer.stop()
        self._switch_visible_lod()
        self.status_message.emit(f"Removed layer: {layer.get('name', layer_id)}")

    def set_layer_visibility(self, layer_id: str, visible: bool):
        layer = self.loaded_layers.get(layer_id)

        if layer is None:
            return

        layer['visible'] = visible

        items = list(layer.get("items", []))
        items.extend(
            item
            for item in layer.get("lod_items", {}).values()
            if item is not None
        )

        geometry_item = layer.get("geometry_item")
        if geometry_item is not None:
            items.append(geometry_item)

        for item in set(items):
            item.setVisible(visible)

        if visible:
            self._switch_visible_lod()


    # ------------------------------------------------------------------
    # Stage 4: navigation helpers
    # ------------------------------------------------------------------

    def fit_all_layers(self):
        """Fit the view to all currently loaded layers."""
        if not self.loaded_layers:
            self.status_message.emit("No layers loaded to fit.")
            return

        bounds = []
        for layer in self.loaded_layers.values():
            gdf = layer.get("gdf")
            if gdf is None:
                continue
            try:
                b = gdf.total_bounds
                if len(b) == 4 and np.all(np.isfinite(b)) and b[2] >= b[0] and b[3] >= b[1]:
                    bounds.append(b)
            except Exception:
                continue

        if not bounds:
            # Point-only layers do not keep a GeoDataFrame in the layer
            # metadata, so derive the range from their graphics instead.
            rect = None
            for layer in self.loaded_layers.values():
                for item in layer.get("items", []):
                    try:
                        r = item.boundingRect()
                    except Exception:
                        continue
                    rect = r if rect is None else rect.united(r)
            if rect is None or not rect.isValid():
                self.status_message.emit("Unable to determine layer extent.")
                return
            self.setRange(xRange=(rect.left(), rect.right()),
                          yRange=(rect.top(), rect.bottom()),
                          padding=0.08)
        else:
            arr = np.asarray(bounds, dtype=float)
            xmin = float(np.min(arr[:, 0]))
            ymin = float(np.min(arr[:, 1]))
            xmax = float(np.max(arr[:, 2]))
            ymax = float(np.max(arr[:, 3]))

            if xmax <= xmin:
                xmax = xmin + 1.0
            if ymax <= ymin:
                ymax = ymin + 1.0

            self.setRange(
                xRange=(xmin, xmax),
                yRange=(ymin, ymax),
                padding=0.08,
            )

        self.status_message.emit("View fitted to all active layers.")

    def zoom_to_layer(self, layer_id: str):
        """Zoom the canvas to one layer."""
        layer = self.loaded_layers.get(layer_id)
        if layer is None:
            return

        gdf = layer.get("gdf")
        if gdf is not None:
            try:
                b = gdf.total_bounds
                if len(b) == 4 and np.all(np.isfinite(b)):
                    xmin, ymin, xmax, ymax = map(float, b)
                    if xmax <= xmin:
                        xmax = xmin + 1.0
                    if ymax <= ymin:
                        ymax = ymin + 1.0
                    self.setRange(
                        xRange=(xmin, xmax),
                        yRange=(ymin, ymax),
                        padding=0.08,
                    )
                    self.status_message.emit(f"Zoomed to {layer['name']}.")
                    return
            except Exception:
                pass

        rect = None
        for item in layer.get("items", []):
            try:
                r = item.boundingRect()
            except Exception:
                continue
            rect = r if rect is None else rect.united(r)

        if rect is not None and rect.isValid():
            self.setRange(
                xRange=(rect.left(), rect.right()),
                yRange=(rect.top(), rect.bottom()),
                padding=0.08,
            )
            self.status_message.emit(f"Zoomed to {layer['name']}.")

    def _mouse_moved(self, scene_pos):
        """Update map coordinates without flooding the status bar."""
        if not self.sceneBoundingRect().contains(scene_pos):
            return

        # QGraphicsScene can emit mouse-move events at very high frequency.
        # Throttle UI updates while keeping the cursor readout responsive.
        now = time.perf_counter()
        if now - self._last_mouse_status_time < 0.05:
            return

        try:
            point = self.getViewBox().mapSceneToView(scene_pos)
            x = float(point.x())
            y = float(point.y())
            if np.isfinite(x) and np.isfinite(y):
                message = f"X: {x:,.3f}    Y: {y:,.3f}"
                if message != self._last_mouse_status_text:
                    self.status_message.emit(message)
                    self._last_mouse_status_text = message
                self._last_mouse_status_time = now
        except Exception:
            pass

    def clear_all_layers(self):
        self._lod_generation += 1
        self._qc_generation += 1
        self._lod_switch_timer.stop()
        self._clear_feature_selection()

        # Clearing the scene removes graphics items. Worker results are
        # ignored after generation changes.
        self.clear()

        self.showGrid(x=True, y=True, alpha=0.3)
        self.loaded_layers.clear()
        self.color_index = 0

        self.layers_cleared.emit()
        self.status_message.emit(
            "Cleared all QC map layers."
        )



from pipeline_route_creator_table import (
    _AttributeTableModel,
    _AttributeFilterProxyModel,
)

class PipelineRouteCreator(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Pipeline Route Shapefile Generator & Fast QC")
        self.resize(1080, 680)
        self.setStyleSheet(bouquet_studio_stylesheet())

        # Stage 4 navigation toolbar.
        navigation_toolbar = QToolBar("Map Navigation", self)
        navigation_toolbar.setMovable(False)
        navigation_toolbar.setToolButtonStyle(Qt.ToolButtonTextOnly)
        self.addToolBar(Qt.TopToolBarArea, navigation_toolbar)

        fit_action = QAction("Fit All", self)
        fit_action.setToolTip("Fit the map to all loaded layers")
        fit_action.triggered.connect(lambda: self.map_canvas.fit_all_layers())
        navigation_toolbar.addAction(fit_action)

        reset_action = QAction("Reset View", self)
        reset_action.setToolTip("Reset the map view")
        reset_action.triggered.connect(lambda: self.map_canvas.enableAutoRange())
        navigation_toolbar.addAction(reset_action)

        navigation_toolbar.addSeparator()

        zoom_in_action = QAction("Zoom In", self)
        zoom_in_action.triggered.connect(lambda: self.map_canvas.getViewBox().scaleBy((0.5, 0.5)))
        navigation_toolbar.addAction(zoom_in_action)

        zoom_out_action = QAction("Zoom Out", self)
        zoom_out_action.triggered.connect(lambda: self.map_canvas.getViewBox().scaleBy((2.0, 2.0)))
        navigation_toolbar.addAction(zoom_out_action)

        navigation_toolbar.addSeparator()

        report_action = QAction("QC Report", self)
        report_action.setToolTip("Export a QC report for all loaded layers")
        report_action.triggered.connect(lambda: self.map_canvas.export_qc_report())
        navigation_toolbar.addAction(report_action)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        layout = QHBoxLayout(central_widget)

        self.splitter = QSplitter(Qt.Horizontal)
        layout.addWidget(self.splitter)

        # Left Control Panel
        left_panel = QWidget()
        left_layout = QVBoxLayout(left_panel)

        left_layout.addWidget(QLabel("Paste Coordinates (Easting Northing):"))
        self.text_edit = QPlainTextEdit()
        self.text_edit.setPlaceholderText("Paste columns here...\n612345.0  5845000.0\n612350.0  5845010.0")
        left_layout.addWidget(self.text_edit)

        # Geometry Type Dropdown
        self.geom_combo = QComboBox()
        self.geom_combo.addItem("Points")
        self.geom_combo.addItem("Line (LineString)")
        self.geom_combo.addItem("Polygon")
        left_layout.addWidget(QLabel("Export Geometry Type:"))
        left_layout.addWidget(self.geom_combo)

        # CRS Dropdown
        self.crs_combo = QComboBox()
        self.crs_combo.addItem("ED50 / UTM 31N (EPSG:23031)", "EPSG:23031")
        self.crs_combo.addItem("ED50 / UTM 32N (EPSG:23032)", "EPSG:23032")
        self.crs_combo.addItem("WGS 84 (EPSG:4326)", "EPSG:4326")
        left_layout.addWidget(QLabel("Select Active Coordinate System:"))
        left_layout.addWidget(self.crs_combo)

        self.export_btn = QPushButton("Generate Shapefile")
        self.export_btn.clicked.connect(self.generate_shapefile)
        left_layout.addWidget(self.export_btn)

        # Stage 15: interactive Create Mode controls.
        create_row = QHBoxLayout()
        self.create_mode_btn = QPushButton("Create Mode")
        self.create_mode_btn.setCheckable(True)
        self.create_mode_btn.setToolTip("Click points directly on the QC map to create a new shapefile")
        self.create_mode_btn.toggled.connect(self.toggle_create_mode)
        create_row.addWidget(self.create_mode_btn)

        self.finish_create_btn = QPushButton("Finish & Generate")
        self.finish_create_btn.setEnabled(False)
        self.finish_create_btn.setToolTip("Finish the clicked geometry and save it as a shapefile")
        self.finish_create_btn.clicked.connect(self.finish_create_mode)
        create_row.addWidget(self.finish_create_btn)
        left_layout.addLayout(create_row)

        self.cancel_create_btn = QPushButton("Cancel Create")
        self.cancel_create_btn.setEnabled(False)
        self.cancel_create_btn.clicked.connect(self.cancel_create_mode)
        left_layout.addWidget(self.cancel_create_btn)

        self.create_points_label = QLabel("Create points: 0")
        left_layout.addWidget(self.create_points_label)

        # Stage 10: batch loading controls.
        load_row = QHBoxLayout()
        self.load_files_btn = QPushButton("Load Shapefiles...")
        self.load_files_btn.setToolTip("Select one or more .shp files to load")
        self.load_files_btn.clicked.connect(self.load_shapefiles_from_dialog)
        load_row.addWidget(self.load_files_btn)

        self.load_folder_btn = QPushButton("Load Folder...")
        self.load_folder_btn.setToolTip("Recursively find and load .shp files in a folder")
        self.load_folder_btn.clicked.connect(self.load_shapefiles_from_folder)
        load_row.addWidget(self.load_folder_btn)
        left_layout.addLayout(load_row)

        # Layer Management QListWidget
        left_layout.addWidget(QLabel("Active Layers (Right-click to change color):"))
        self.layer_list = QListWidget()
        self.layer_list.itemChanged.connect(self.on_layer_item_changed)

        self.layer_list.setContextMenuPolicy(Qt.CustomContextMenu)
        self.layer_list.customContextMenuRequested.connect(self.show_layer_context_menu)

        left_layout.addWidget(self.layer_list)

        self.clear_map_btn = QPushButton("Clear QC Canvas")
        self.clear_map_btn.clicked.connect(self.clear_canvas)
        left_layout.addWidget(self.clear_map_btn)

        # Right Panel: pyqtgraph Canvas
        self.map_canvas = QCMapCanvas(self)

        self.splitter.addWidget(left_panel)
        self.splitter.addWidget(self.map_canvas)
        self.splitter.setSizes([380, 700])

        # Signal connections
        self.statusBar().showMessage("Ready. Paste coordinates or drag-and-drop multiple .shp files.")
        self.map_canvas.status_message.connect(self.statusBar().showMessage)
        self.map_canvas.layer_added.connect(self.add_layer_to_list)
        self.map_canvas.layer_qc_updated.connect(self.update_layer_qc_status)
        self.map_canvas.layers_cleared.connect(self.layer_list.clear)
        self.map_canvas.create_points_changed.connect(self.update_create_point_count)
        self.map_canvas.create_finished.connect(self.create_shapefile_from_points)

        # Stage 10: lightweight keyboard workflow.
        QShortcut(Qt.CTRL | Qt.Key_F, self, activated=self.map_canvas.fit_all_layers)
        QShortcut(Qt.CTRL | Qt.Key_R, self, activated=self.map_canvas.enableAutoRange)
        QShortcut(Qt.CTRL | Qt.Key_L, self, activated=self.clear_canvas)
        QShortcut(Qt.Key_Delete, self.layer_list, activated=self.remove_selected_layer)

        self._create_application_menus()
        self._restore_window_state()

    def _create_application_menus(self):
        """Create a small, predictable application menu for the final workflow."""
        file_menu = self.menuBar().addMenu("File")

        load_files_action = QAction("Load Shapefiles...", self)
        load_files_action.setShortcut("Ctrl+O")
        load_files_action.triggered.connect(self.load_shapefiles_from_dialog)
        file_menu.addAction(load_files_action)

        load_folder_action = QAction("Load Folder...", self)
        load_folder_action.triggered.connect(self.load_shapefiles_from_folder)
        file_menu.addAction(load_folder_action)

        file_menu.addSeparator()
        report_action = QAction("QC Report...", self)
        report_action.triggered.connect(lambda: self.map_canvas.export_qc_report())
        file_menu.addAction(report_action)

        file_menu.addSeparator()
        exit_action = QAction("Exit", self)
        exit_action.setShortcut("Ctrl+Q")
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)

        view_menu = self.menuBar().addMenu("View")
        fit_action = QAction("Fit All", self)
        fit_action.setShortcut("Ctrl+F")
        fit_action.triggered.connect(self.map_canvas.fit_all_layers)
        view_menu.addAction(fit_action)

        reset_action = QAction("Reset View", self)
        reset_action.setShortcut("Ctrl+R")
        reset_action.triggered.connect(self.map_canvas.enableAutoRange)
        view_menu.addAction(reset_action)

        layers_menu = self.menuBar().addMenu("Layers")
        show_all_action = QAction("Show All", self)
        show_all_action.triggered.connect(self.show_all_layers)
        layers_menu.addAction(show_all_action)

        hide_all_action = QAction("Hide All", self)
        hide_all_action.triggered.connect(self.hide_all_layers)
        layers_menu.addAction(hide_all_action)

        layers_menu.addSeparator()
        remove_action = QAction("Remove Selected", self)
        remove_action.setShortcut("Delete")
        remove_action.triggered.connect(self.remove_selected_layer)
        layers_menu.addAction(remove_action)

        help_menu = self.menuBar().addMenu("Help")
        shortcuts_action = QAction("Keyboard Shortcuts", self)
        shortcuts_action.triggered.connect(self.show_keyboard_shortcuts)
        help_menu.addAction(shortcuts_action)

        about_action = QAction("About", self)
        about_action.triggered.connect(self.show_about)
        help_menu.addAction(about_action)

    def show_all_layers(self):
        self.layer_list.blockSignals(True)
        try:
            for row in range(self.layer_list.count()):
                self.layer_list.item(row).setCheckState(Qt.Checked)
        finally:
            self.layer_list.blockSignals(False)
        for layer_id in self.map_canvas.loaded_layers:
            self.map_canvas.set_layer_visibility(layer_id, True)

    def hide_all_layers(self):
        self.layer_list.blockSignals(True)
        try:
            for row in range(self.layer_list.count()):
                self.layer_list.item(row).setCheckState(Qt.Unchecked)
        finally:
            self.layer_list.blockSignals(False)
        for layer_id in self.map_canvas.loaded_layers:
            self.map_canvas.set_layer_visibility(layer_id, False)

    def show_keyboard_shortcuts(self):
        QMessageBox.information(
            self,
            "Keyboard Shortcuts",
            "Ctrl+O  Load shapefiles\n"
            "Ctrl+F  Fit all layers\n"
            "Ctrl+R  Reset view\n"
            "Ctrl+L  Clear canvas\n"
            "Delete  Remove selected layer\n"
            "Create Mode  Click points on the map to create a shapefile\n"
            "Ctrl+Q  Exit",
        )

    def show_about(self):
        QMessageBox.about(
            self,
            "About Pipeline Route Shapefile Generator & Fast QC",
            "Pipeline Route Shapefile Generator & Fast QC\n\n"
            "A fast desktop tool for creating, visualising and quality-checking shapefiles.",
        )

    def _restore_window_state(self):
        settings = QSettings("42 Decisions", "PipelineRouteCreator")
        geometry = settings.value("window_geometry")
        splitter_state = settings.value("splitter_state")
        if geometry:
            self.restoreGeometry(geometry)
        if splitter_state:
            self.splitter.restoreState(splitter_state)

    def closeEvent(self, event):
        settings = QSettings("42 Decisions", "PipelineRouteCreator")
        settings.setValue("window_geometry", self.saveGeometry())
        settings.setValue("splitter_state", self.splitter.saveState())
        settings.sync()
        super().closeEvent(event)

    def load_shapefiles_from_dialog(self):
        files, _ = QFileDialog.getOpenFileNames(
            self,
            "Load Shapefiles",
            "",
            "Shapefiles (*.shp)",
        )
        if files:
            self.load_shapefiles(files)

    def load_shapefiles_from_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "Select Shapefile Folder")
        if not folder:
            return

        root = Path(folder)
        files = sorted(
            str(path) for path in root.rglob("*.shp")
            if path.is_file()
        )
        if not files:
            QMessageBox.information(
                self, "Load Folder", "No .shp files were found in the selected folder."
            )
            return

        self.load_shapefiles(files)

    def load_shapefiles(self, files):
        """Load a batch of shapefiles with progress feedback."""
        unique_files = []
        seen = set()
        for filepath in files:
            try:
                key = str(Path(filepath).resolve()).lower()
            except Exception:
                key = str(filepath).lower()
            if key not in seen:
                seen.add(key)
                unique_files.append(str(filepath))

        total = len(unique_files)
        loaded = 0
        skipped = 0

        self.load_files_btn.setEnabled(False)
        self.load_folder_btn.setEnabled(False)
        try:
            for number, filepath in enumerate(unique_files, start=1):
                name = Path(filepath).name
                self.statusBar().showMessage(
                    f"Loading {number}/{total}: {name}"
                )
                QApplication.processEvents()
                if self.map_canvas.add_shapefile_layer(filepath):
                    loaded += 1
                else:
                    skipped += 1
                QApplication.processEvents()
        finally:
            self.load_files_btn.setEnabled(True)
            self.load_folder_btn.setEnabled(True)

        self.statusBar().showMessage(
            f"Batch load complete: {loaded} loaded, {skipped} skipped. "
            f"Active layers: {len(self.map_canvas.loaded_layers)}."
        )

    def remove_selected_layer(self):
        item = self.layer_list.currentItem()
        if item is None:
            return
        layer_id = item.data(Qt.UserRole)
        if layer_id:
            self.remove_layer(layer_id)

    def get_active_crs(self) -> str:
        return self.crs_combo.currentData()

    def add_layer_to_list(self, layer_id: str, display_name: str, color_hex: str):
        item = QListWidgetItem(display_name)
        item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
        item.setCheckState(Qt.Checked)
        item.setData(Qt.UserRole, layer_id)
        item.setData(Qt.UserRole + 1, display_name)
        item.setData(Qt.UserRole + 2, color_hex)
        item.setIcon(make_status_icon("CHECKING..."))
        item.setToolTip("Geometry QC: CHECKING...")
        self.layer_list.addItem(item)

    def update_layer_qc_status(self, layer_id: str, status: str):
        for index in range(self.layer_list.count()):
            item = self.layer_list.item(index)
            if item.data(Qt.UserRole) != layer_id:
                continue
            layer = self.map_canvas.loaded_layers.get(layer_id, {})
            name = layer.get("name", item.data(Qt.UserRole + 1) or item.text())
            color_hex = layer.get("color", item.data(Qt.UserRole + 2) or "#808080")
            item.setText(name)
            item.setIcon(make_status_icon(status))
            item.setToolTip(f"Geometry QC: {status}\nLayer color: {color_hex}")
            return

    def on_layer_item_changed(self, item: QListWidgetItem):
        layer_id = item.data(Qt.UserRole)
        if layer_id:
            is_visible = item.checkState() == Qt.Checked
            self.map_canvas.set_layer_visibility(layer_id, is_visible)

    def show_layer_context_menu(self, pos):
        item = self.layer_list.itemAt(pos)
        if not item:
            return

        menu = QMenu(self)
        zoom_action = menu.addAction("Zoom to Layer")
        info_action = menu.addAction("Layer Information")
        qc_action = menu.addAction("Geometry QC")
        attribute_action = menu.addAction("Attribute Table")
        clear_selection_action = menu.addAction("Clear Feature Selection")
        change_color_action = menu.addAction("Change Color...")
        menu.addSeparator()
        remove_action = menu.addAction("Remove Layer")

        action = menu.exec(self.layer_list.mapToGlobal(pos))

        if action == zoom_action:
            layer_id = item.data(Qt.UserRole)
            if layer_id:
                self.map_canvas.zoom_to_layer(layer_id)
        elif action == info_action:
            layer_id = item.data(Qt.UserRole)
            if layer_id:
                self.map_canvas.show_layer_information(layer_id)
        elif action == qc_action:
            layer_id = item.data(Qt.UserRole)
            if layer_id:
                self.map_canvas.show_geometry_qc(layer_id)
        elif action == attribute_action:
            layer_id = item.data(Qt.UserRole)
            if layer_id:
                self.map_canvas.show_attribute_table(layer_id)
        elif action == clear_selection_action:
            self.map_canvas._clear_feature_selection()
        elif action == change_color_action:
            self.change_item_color(item)
        elif action == remove_action:
            layer_id = item.data(Qt.UserRole)
            if layer_id:
                self.remove_layer(layer_id)

    def remove_layer(self, layer_id: str):
        """Remove one layer without disturbing the other loaded layers."""
        layer = self.map_canvas.loaded_layers.get(layer_id)
        if layer is None:
            return

        name = layer.get("name", layer_id)
        self.map_canvas.remove_layer(layer_id)

        for row in range(self.layer_list.count() - 1, -1, -1):
            item = self.layer_list.item(row)
            if item.data(Qt.UserRole) == layer_id:
                self.layer_list.takeItem(row)
                del item
                break

        self.statusBar().showMessage(f"Removed layer: {name}")

    def change_item_color(self, item: QListWidgetItem):
        layer_id = item.data(Qt.UserRole)
        if not layer_id:
            return

        color = QColorDialog.getColor(parent=self, title="Select Layer Color")
        if color.isValid():
            color_hex = color.name()
            item.setData(Qt.UserRole + 2, color_hex)
            self.map_canvas.change_layer_color(layer_id, color_hex)
            layer = self.map_canvas.loaded_layers.get(layer_id, {})
            qc_status = layer.get("qc_status", "CHECKING...")
            item.setToolTip(f"Geometry QC: {qc_status}\nLayer color: {color_hex}")

    def toggle_create_mode(self, enabled):
        if enabled:
            geometry_type = self.geom_combo.currentText()
            self.map_canvas.set_create_mode(True, geometry_type)
            self.finish_create_btn.setEnabled(True)
            self.cancel_create_btn.setEnabled(True)
            self.export_btn.setEnabled(False)
            self.create_mode_btn.setText("Create Mode: ON")
            self.geom_combo.setEnabled(False)
            self.crs_combo.setEnabled(False)
        else:
            self.map_canvas.cancel_create_mode()
            self.finish_create_btn.setEnabled(False)
            self.cancel_create_btn.setEnabled(False)
            self.export_btn.setEnabled(True)
            self.create_mode_btn.setText("Create Mode")
            self.geom_combo.setEnabled(True)
            self.crs_combo.setEnabled(True)

    def finish_create_mode(self):
        points = self.map_canvas.creation_points()
        geom_type = self.geom_combo.currentText()
        if geom_type == "Points":
            minimum = 1
        elif geom_type == "Line (LineString)":
            minimum = 2
        else:
            minimum = 3

        if len(points) < minimum:
            self.statusBar().showMessage(
                f"Error: {geom_type} requires at least {minimum} clicked point(s)."
            )
            return

        self.map_canvas.finish_create_mode()

    def cancel_create_mode(self):
        self.create_mode_btn.blockSignals(True)
        self.create_mode_btn.setChecked(False)
        self.create_mode_btn.blockSignals(False)
        self.map_canvas.cancel_create_mode()
        self.finish_create_btn.setEnabled(False)
        self.cancel_create_btn.setEnabled(False)
        self.export_btn.setEnabled(True)
        self.create_mode_btn.setText("Create Mode")
        self.geom_combo.setEnabled(True)
        self.crs_combo.setEnabled(True)
        self.statusBar().showMessage("Create Mode cancelled.")

    def update_create_point_count(self, count):
        self.create_points_label.setText(f"Create points: {count}")

    def create_shapefile_from_points(self, points):
        # Restore the normal controls before validation/save.
        self.create_mode_btn.blockSignals(True)
        self.create_mode_btn.setChecked(False)
        self.create_mode_btn.blockSignals(False)
        self.finish_create_btn.setEnabled(False)
        self.cancel_create_btn.setEnabled(False)
        self.export_btn.setEnabled(True)
        self.create_mode_btn.setText("Create Mode")
        self.geom_combo.setEnabled(True)
        self.crs_combo.setEnabled(True)
        self.create_points_label.setText(f"Create points: {len(points)}")

        if not points:
            self.statusBar().showMessage("Create Mode finished without any points.")
            return

        selected_crs = self.get_active_crs()
        geom_type = self.geom_combo.currentText()

        if geom_type == "Points":
            geometries = [Point(x, y) for x, y in points]
            gdf = gpd.GeoDataFrame({'geometry': geometries}, crs=selected_crs)
        elif geom_type == "Line (LineString)":
            if len(points) < 2:
                self.statusBar().showMessage("Error: Need at least 2 clicked points to create a Line.")
                return
            line_geom = LineString(points)
            gdf = gpd.GeoDataFrame({'geometry': [line_geom]}, crs=selected_crs)
        else:
            distinct_points = list(dict.fromkeys(points))
            if len(distinct_points) < 3:
                self.statusBar().showMessage(
                    "Error: Need at least 3 distinct clicked points to create a Polygon."
                )
                return
            polygon_geom = Polygon(points)
            if polygon_geom.is_empty or polygon_geom.area <= 0:
                self.statusBar().showMessage(
                    "Error: Polygon has zero area. Check the clicked point order."
                )
                return
            if not polygon_geom.is_valid:
                try:
                    reason = shapely.is_valid_reason(polygon_geom)
                except Exception:
                    reason = "Polygon geometry is invalid."
                QMessageBox.warning(
                    self,
                    "Invalid Polygon",
                    f"The clicked points do not form a valid polygon.\n\n"
                    f"Reason: {reason}\n\n"
                    "Check the point order and make sure the boundary does not self-intersect.",
                )
                self.statusBar().showMessage("Polygon creation cancelled: invalid geometry.")
                return
            gdf = gpd.GeoDataFrame({'geometry': [polygon_geom]}, crs=selected_crs)

        save_path, _ = QFileDialog.getSaveFileName(
            self, "Save Created Shapefile", "", "Shapefiles (*.shp)"
        )
        if not save_path:
            self.statusBar().showMessage("Create Mode finished; export cancelled.")
            return

        try:
            gdf.to_file(save_path)
            self.map_canvas.add_shapefile_layer(save_path, target_crs=selected_crs)
            filename = Path(save_path).name
            self.statusBar().showMessage(
                f"Created {geom_type} from {len(points)} clicked point(s): "
                f"{filename} ({selected_crs})."
            )
        except Exception as e:
            self.statusBar().showMessage(f"Export failed: {str(e)}")

    def generate_shapefile(self):
        raw_text = self.text_edit.toPlainText().strip()
        if not raw_text:
            self.statusBar().showMessage("Error: No coordinate input found.")
            return

        points = []
        for line in raw_text.split('\n'):
            parts = line.strip().split()
            if len(parts) >= 2:
                try:
                    easting = float(parts[0])
                    northing = float(parts[1])
                    points.append((easting, northing))
                except ValueError:
                    continue

        if not points:
            self.statusBar().showMessage("Error: No valid numerical coordinate pairs parsed.")
            return

        selected_crs = self.get_active_crs()
        geom_type = self.geom_combo.currentText()

        if geom_type == "Points":
            geometries = [Point(x, y) for x, y in points]
            gdf = gpd.GeoDataFrame({'geometry': geometries}, crs=selected_crs)
        elif geom_type == "Line (LineString)":
            if len(points) < 2:
                self.statusBar().showMessage("Error: Need at least 2 coordinate pairs to create a Line.")
                return
            line_geom = LineString(points)
            gdf = gpd.GeoDataFrame({'geometry': [line_geom]}, crs=selected_crs)
        else:
            # Polygon boundaries are defined by the supplied points in order.
            # The user does not need to repeat the first point; Shapely closes
            # the ring automatically.
            distinct_points = list(dict.fromkeys(points))
            if len(distinct_points) < 3:
                self.statusBar().showMessage(
                    "Error: Need at least 3 distinct coordinate pairs to create a Polygon."
                )
                return

            polygon_geom = Polygon(points)

            if polygon_geom.is_empty or polygon_geom.area <= 0:
                self.statusBar().showMessage(
                    "Error: Polygon has zero area. Check the coordinate order."
                )
                return

            if not polygon_geom.is_valid:
                try:
                    reason = shapely.is_valid_reason(polygon_geom)
                except Exception:
                    reason = "Polygon geometry is invalid."
                QMessageBox.warning(
                    self,
                    "Invalid Polygon",
                    f"The supplied coordinates do not form a valid polygon.\n\n"
                    f"Reason: {reason}\n\n"
                    "Check the point order and make sure the boundary does not self-intersect."
                )
                self.statusBar().showMessage("Polygon creation cancelled: invalid geometry.")
                return

            gdf = gpd.GeoDataFrame({'geometry': [polygon_geom]}, crs=selected_crs)
            self.statusBar().showMessage(
                f"Polygon ready: area {polygon_geom.area:,.3f}, "
                f"perimeter {polygon_geom.length:,.3f}."
            )

        save_path, _ = QFileDialog.getSaveFileName(
            self, "Save Shapefile", "", "Shapefiles (*.shp)"
        )

        if save_path:
            try:
                gdf.to_file(save_path)
                self.map_canvas.add_shapefile_layer(save_path, target_crs=selected_crs)
                filename = Path(save_path).name
                if geom_type == "Polygon":
                    polygon_geom = gdf.geometry.iloc[0]
                    self.statusBar().showMessage(
                        f"Exported Polygon to {filename} ({selected_crs}) — "
                        f"area {polygon_geom.area:,.3f}, perimeter {polygon_geom.length:,.3f}."
                    )
                else:
                    self.statusBar().showMessage(
                        f"Exported {geom_type} to {filename} ({selected_crs})."
                    )
            except Exception as e:
                self.statusBar().showMessage(f"Export failed: {str(e)}")
        else:
            self.statusBar().showMessage("Export cancelled.")

    def clear_canvas(self):
        self.map_canvas.clear_all_layers()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = PipelineRouteCreator()
    window.show()
    sys.exit(app.exec())